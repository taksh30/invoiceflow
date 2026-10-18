"""Unit tests for monthly CSV and Excel export generation."""

from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from invoiceflow.db.tables import InvoiceRow
from invoiceflow.export.csv_xlsx import (
    export_monthly_data,
    export_to_csv,
    export_to_excel,
    prepare_export_dataframes,
)


def _make_mock_invoice(
    inv_id: int,
    vendor_name: str,
    total: str,
    invoice_date: date = date(2024, 5, 15),
    subtotal: str = "1000.00",
    tax_total: str = "180.00",
    cgst: str = "90.00",
    sgst: str = "90.00",
) -> InvoiceRow:
    """Helper to create dummy InvoiceRow instances for testing exports."""
    inv = InvoiceRow(
        document_id=inv_id,
        vendor_name=vendor_name,
        vendor_tax_id="27AAACR1234F1Z5",
        invoice_number=f"INV-2024-{inv_id:03d}",
        invoice_date=invoice_date,
        currency="INR",
        subtotal=Decimal(subtotal),
        cgst=Decimal(cgst),
        sgst=Decimal(sgst),
        igst=Decimal("0.00"),
        tax_total=Decimal(tax_total),
        total=Decimal(total),
        overall_confidence=0.95,
    )
    inv.id = inv_id
    return inv


def test_prepare_export_dataframes_empty() -> None:
    """Verify empty input returns empty DataFrames with correct column schema."""
    df_inv, df_summary = prepare_export_dataframes([])

    assert df_inv.empty
    assert "Invoice ID" in df_inv.columns
    assert "Vendor Name" in df_inv.columns
    assert "Total" in df_inv.columns

    assert df_summary.empty
    assert "Vendor Name" in df_summary.columns
    assert "Invoice Count" in df_summary.columns


def test_prepare_export_dataframes_aggregation() -> None:
    """Verify invoice details and vendor summary aggregation calculations."""
    invoices = [
        _make_mock_invoice(1, "Acme Corp", "1180.00", subtotal="1000.00", tax_total="180.00"),
        _make_mock_invoice(2, "Acme Corp", "2360.00", subtotal="2000.00", tax_total="360.00"),
        _make_mock_invoice(3, "Beta Ltd", "590.00", subtotal="500.00", tax_total="90.00"),
    ]

    df_inv, df_summary = prepare_export_dataframes(invoices)

    assert len(df_inv) == 3
    assert list(df_inv["Vendor Name"]) == ["Acme Corp", "Acme Corp", "Beta Ltd"]

    # Check vendor summary
    assert len(df_summary) == 2
    acme_row = df_summary[df_summary["Vendor Name"] == "Acme Corp"].iloc[0]
    assert acme_row["Invoice Count"] == 2
    assert acme_row["Subtotal"] == pytest.approx(3000.0)
    assert acme_row["Tax Total"] == pytest.approx(540.0)
    assert acme_row["Total"] == pytest.approx(3540.0)

    beta_row = df_summary[df_summary["Vendor Name"] == "Beta Ltd"].iloc[0]
    assert beta_row["Invoice Count"] == 1
    assert beta_row["Total"] == pytest.approx(590.0)


def test_export_to_excel_creates_two_sheets(tmp_path: Path) -> None:
    """Verify Excel export creates a valid workbook containing 'Invoices' and 'Vendor Summary'."""
    invoices = [
        _make_mock_invoice(1, "Acme Corp", "1180.00"),
        _make_mock_invoice(2, "Beta Ltd", "590.00"),
    ]
    out_file = tmp_path / "invoices_2024_05.xlsx"

    res_path = export_to_excel(invoices, out_file)
    assert res_path.exists()

    # Load with openpyxl to inspect sheet names directly
    wb = openpyxl.load_workbook(res_path)
    assert "Invoices" in wb.sheetnames
    assert "Vendor Summary" in wb.sheetnames

    # Load with pandas to verify content
    sheets = pd.read_excel(res_path, sheet_name=None)
    assert len(sheets["Invoices"]) == 2
    assert len(sheets["Vendor Summary"]) == 2


def test_export_to_csv_creates_files(tmp_path: Path) -> None:
    """Verify CSV export produces detailed invoices CSV and companion summary CSV."""
    invoices = [
        _make_mock_invoice(1, "Acme Corp", "1180.00"),
        _make_mock_invoice(2, "Beta Ltd", "590.00"),
    ]
    out_file = tmp_path / "invoices_2024_05.csv"

    inv_path, summary_path = export_to_csv(invoices, out_file)
    assert inv_path.exists()
    assert summary_path.exists()
    assert summary_path.name == "invoices_2024_05_summary.csv"

    df_read = pd.read_csv(inv_path)
    assert len(df_read) == 2
    assert "Acme Corp" in df_read["Vendor Name"].values


def test_export_monthly_data_dispatch(tmp_path: Path) -> None:
    """Verify format dispatcher correctly invokes xlsx or csv export."""
    invoices = [_make_mock_invoice(1, "Acme Corp", "1180.00")]

    xlsx_path = export_monthly_data(invoices, tmp_path / "test.xlsx", export_format="xlsx")
    assert xlsx_path.suffix == ".xlsx"
    assert xlsx_path.exists()

    csv_path = export_monthly_data(invoices, tmp_path / "test.csv", export_format="csv")
    assert csv_path.suffix == ".csv"
    assert csv_path.exists()

    with pytest.raises(ValueError, match="Unsupported export format"):
        export_monthly_data(invoices, tmp_path / "test.pdf", export_format="pdf")  # type: ignore[arg-type]
